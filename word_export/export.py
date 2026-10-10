"""Export saved OCR results and ordered PNG/JPEG pages to an editable Word file."""

from __future__ import annotations

import os
import tempfile
import unicodedata
from collections.abc import Sequence
from contextlib import contextmanager
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from statistics import median
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree
from PIL import Image

from word_export.content import convert_regions, dotted_list
from word_export.measure import (
    FontMetrics,
    cell_alignment,
    cell_appearance,
    flat_background,
    missing_graphics,
    source_font_size,
    table_geometry,
    text_color,
    text_lines,
    text_rotation,
    word_tokens,
)
from word_export.pages import Page, Region, load_image
from word_export.scan import content_box, estimate_skew, rotate_scan, transform_box

WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
EMU_PER_POINT = 12700


@contextmanager
def prepared_page(page: Page, warnings: list[str], *, deskew: bool = True):
    """Use one coordinate transform for the source pixels and saved OCR regions."""
    with load_image(page.path) as image:
        if image.size != (page.width, page.height):
            raise ValueError(f"{page.path}: source image changed during export")
        # Exclude photographs and intentionally quarter-turned regions from
        # page-wide tilt estimation. They are still transformed with the page.
        with Image.new("RGB", image.size, "white") as evidence:
            for region in page.regions:
                if region.category not in {"Picture", "Formula"} and not region.rotation:
                    with image.crop(region.bbox) as crop:
                        evidence.paste(crop, region.bbox[:2])
            estimate = estimate_skew(evidence) if deskew else None
        if estimate is None or not estimate.confident or not estimate.angle:
            if estimate is not None and not estimate.confident and abs(estimate.angle) >= 0.15:
                warnings.append(f"{page.path.name}: uncertain scan skew ({estimate.reason}); retained original coordinates")
            yield page, image
            return
        corrected, matrix = rotate_scan(image, estimate.angle)
        with corrected:
            regions = []
            for region in page.regions:
                box = transform_box(region.bbox, matrix, corrected.size)
                if region.category not in {"Picture", "Formula"} and not region.rotation:
                    box = content_box(corrected, region.bbox, matrix, table=region.category == "Table")
                regions.append(replace(region, bbox=box))
            warnings.append(f"{page.path.name}: deskewed by {estimate.angle:+.2f} degrees with matching OCR coordinates; review layout (OCR text is unchanged)")
            yield replace(page, width=corrected.width, height=corrected.height, regions=regions), corrected


def page_geometry(page: Page) -> tuple[int, int, float, float, float]:
    width, height = (Mm(297), Mm(210)) if page.width > page.height else (Mm(210), Mm(297))
    scale = min(width / page.width, height / page.height)
    left = (width - page.width * scale) / 2
    top = (height - page.height * scale) / 2
    return width, height, scale, left, top


def region_geometry(page: Page, region: Region) -> tuple[int, int, int, int]:
    _, _, scale, left, top = page_geometry(page)
    x1, y1, x2, y2 = region.bbox
    return round(left + x1 * scale), round(top + y1 * scale), round((x2 - x1) * scale), round((y2 - y1) * scale)


def xml_element(tag: str, **attributes: Any) -> Any:
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(key, str(value))
    return element


def page_anchor(identifier: int, label: str, geometry: tuple[int, int, int, int]) -> Any:
    left, top, width, height = geometry
    anchor = xml_element(
        "wp:anchor", distT=0, distB=0, distL=0, distR=0, simplePos=0,
        relativeHeight=identifier, behindDoc=0, locked=1, layoutInCell=1, allowOverlap=1,
    )
    anchor.append(xml_element("wp:simplePos", x=0, y=0))
    for axis, offset in (("H", left), ("V", top)):
        position = xml_element(f"wp:position{axis}", relativeFrom="page")
        value = xml_element("wp:posOffset")
        value.text = str(offset)
        position.append(value)
        anchor.append(position)
    anchor.append(xml_element("wp:extent", cx=width, cy=height))
    anchor.append(xml_element("wp:effectExtent", l=0, t=0, r=0, b=0))
    anchor.append(xml_element("wp:wrapNone"))
    anchor.append(xml_element("wp:docPr", id=identifier, name=label))
    anchor.append(xml_element("wp:cNvGraphicFramePr"))
    return anchor


def add_textbox(paragraph: Any, identifier: int, region: Region, geometry: tuple[int, int, int, int], nodes: list[Any]) -> None:
    anchor = page_anchor(identifier, region.label, geometry)
    graphic = xml_element("a:graphic")
    data = xml_element("a:graphicData", uri=WPS_NS)
    shape = etree.SubElement(data, f"{{{WPS_NS}}}wsp", nsmap={"wps": WPS_NS})
    etree.SubElement(shape, f"{{{WPS_NS}}}cNvSpPr", txBox="1")
    properties = etree.SubElement(shape, f"{{{WPS_NS}}}spPr")
    transform = xml_element("a:xfrm")
    transform.append(xml_element("a:off", x=0, y=0))
    transform.append(xml_element("a:ext", cx=geometry[2], cy=geometry[3]))
    properties.append(transform)
    rectangle = xml_element("a:prstGeom", prst="rect")
    rectangle.append(xml_element("a:avLst"))
    properties.append(rectangle)
    properties.append(xml_element("a:noFill"))
    line = xml_element("a:ln")
    line.append(xml_element("a:noFill"))
    properties.append(line)
    textbox = etree.SubElement(shape, f"{{{WPS_NS}}}txbx")
    content = xml_element("w:txbxContent")
    content.extend(nodes)
    if not nodes or nodes[-1].tag != qn("w:p"):
        # Word requires a trailing paragraph after a table in a text box.
        trailing = xml_element("w:p")
        Paragraph(trailing, paragraph._parent).paragraph_format.line_spacing = Pt(1)
        content.append(trailing)
    textbox.append(content)
    body = etree.SubElement(
        shape, f"{{{WPS_NS}}}bodyPr", rot="10800000" if region.rotation == 180 else "0",
        vert={90: "vert", 270: "vert270"}.get(region.rotation, "horz"), wrap="square",
        lIns="0", tIns="0", rIns="0", bIns="0", anchor="t",
        vertOverflow="overflow", horzOverflow="overflow",
    )
    body.append(xml_element("a:noAutofit"))
    graphic.append(data)
    anchor.append(graphic)
    drawing = xml_element("w:drawing")
    drawing.append(anchor)
    paragraph.add_run()._r.append(drawing)


def add_picture(paragraph: Any, identifier: int, region: Region, geometry: tuple[int, int, int, int], image: Image.Image,
                overlays: Sequence[Region] = (), warnings: list[str] | None = None) -> None:
    with image.crop(region.bbox) as crop:
        x1, y1, x2, y2 = region.bbox
        for other in overlays:
            if other.category in {"Picture", "Table", "Formula"} or not other.text.strip():
                continue
            a, b, c, d = other.bbox
            if c <= x1 or a >= x2 or d <= y1 or b >= y2:
                continue
            # Examine the complete text region even when only part intersects
            # the picture; a picture crop can contain half a line of text.
            patch = (max(0, a - 3), max(0, b - 3), min(image.width, c + 3), min(image.height, d + 3))
            with image.crop(patch) as sample:
                background = flat_background(sample)
            if background is not None:
                crop.paste(background, (max(a, x1) - x1, max(b, y1) - y1,
                                        min(c, x2) - x1, min(d, y2) - y1))
            if warnings is not None:
                action = "replaced duplicate picture pixels with editable text" if background is not None else "unresolved picture/text overlap; background is not safely replaceable"
                warnings.append(f"{other.label}: {action} ({region.label})")
        buffer = BytesIO()
        crop.save(buffer, format="PNG")
    buffer.seek(0)
    inline = paragraph.add_run().add_picture(buffer, width=geometry[2], height=geometry[3])._inline
    anchor = page_anchor(identifier, region.label, geometry)
    # Keep recognized text visible even when OCR lists its enclosing picture
    # later in reading order. The saved region order itself stays unchanged.
    anchor.set("behindDoc", "1")
    anchor.append(inline.graphic)
    inline.getparent().replace(inline, anchor)


def visible_text(node: Any) -> str:
    pieces = []
    for child in node.iter():
        if child.tag in {qn("w:t"), qn("m:t")}:
            pieces.append(child.text or "")
        elif child.tag in {qn("w:br"), qn("w:tab")}:
            pieces.append("\n" if child.tag == qn("w:br") else " ")
    return "".join(pieces)


def estimated_height(nodes: list[Any], width: float, size: float, metrics: FontMetrics, pitch: float | None = None) -> float:
    """Measure styled glyph advances, reserving extra height for native equations."""
    height = 0.0
    for node in nodes:
        if node.tag == qn("w:tbl"):
            grid = node.findall("./" + qn("w:tblGrid") + "/" + qn("w:gridCol"))
            widths = [int(item.get(qn("w:w"), "1")) for item in grid] or [1]
            for row in node.findall(qn("w:tr")):
                heights = []
                margins = []
                index = 0
                for cell in row.findall(qn("w:tc")):
                    span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
                    count = int(span.get(qn("w:val"))) if span is not None else 1
                    cell_width = width * sum(widths[index:index + count]) / sum(widths)
                    inset = cell.find("./" + qn("w:tcPr") + "/" + qn("w:tcMar"))
                    padding = {}
                    for side in ("left", "right", "top", "bottom"):
                        value = inset.find(qn("w:" + side)) if inset is not None else None
                        padding[side] = int(value.get(qn("w:w"), "0")) / 20 if value is not None else 0
                    heights.append(estimated_height(list(cell), max(1, cell_width - padding['left'] - padding['right']), size, metrics, pitch))
                    margins.append(padding['top'] + padding['bottom'])
                    index += count
                row_height = row.find("./" + qn("w:trPr") + "/" + qn("w:trHeight"))
                minimum = int(row_height.get(qn("w:val"), "0")) / 20 if row_height is not None else 0
                height += max(minimum, max((value + margin for value, margin in zip(heights, margins)), default=metrics.line_height(size)))
        elif node.tag == qn("w:p"):
            spacing = node.find("./" + qn("w:pPr") + "/" + qn("w:spacing"))
            if spacing is not None:
                height += sum(int(spacing.get(qn(name), "0")) / 20 for name in ("w:before", "w:after"))
            line_height = max(pitch or 0, metrics.line_height(size))
            used = 0.0
            current_height = line_height
            indent = node.find("./" + qn("w:pPr") + "/" + qn("w:ind"))
            available = max(1, width - (int(indent.get(qn("w:left"), "0")) / 20 if indent is not None else 0) - 1)
            for child in node:
                if child.tag in {qn("m:oMath"), qn("m:oMathPara")}:
                    text = visible_text(child)
                    advance = metrics.width(text, size) + size * 0.3
                    factor = 2.4 if any(True for _ in child.iter(qn("m:f"), qn("m:nary"))) else 1.6
                    rows = list(child.iter(qn("m:mr")))
                    factor = max(factor, len(rows) * 1.2)
                    tokens = [(advance, line_height * factor, False)]
                else:
                    tokens = []
                    for run in child.iter(qn("w:r")) if child.tag != qn("w:r") else [child]:
                        bold = run.find("./" + qn("w:rPr") + "/" + qn("w:b")) is not None
                        italic = run.find("./" + qn("w:rPr") + "/" + qn("w:i")) is not None
                        for part in run:
                            if part.tag == qn("w:br"):
                                tokens.append((0, line_height, True))
                            elif part.tag == qn("w:t"):
                                tokens.extend((metrics.width(token, size, bold, italic), line_height, False)
                                              for token in word_tokens(part.text or ""))
                for advance, token_height, newline in tokens:
                    if newline or used and used + advance > available:
                        height += current_height
                        used, current_height = 0.0, line_height
                    if not newline:
                        # Very long unbroken tokens wrap in Word too.
                        while advance > available:
                            height += max(current_height, token_height)
                            advance -= available
                        used += advance
                        current_height = max(current_height, token_height)
            height += current_height
    return height


def format_paragraphs(nodes: list[Any], parent: Any, size: float, metrics: FontMetrics,
                      gaps: dict[int, float], pitch: float | None) -> None:
    """Apply the same indentation and leading used by the fit calculation."""
    for node in nodes:
        for paragraph in node.iter(qn("w:p")):
            wrapper = Paragraph(paragraph, parent)
            formatting = wrapper.paragraph_format
            formatting.space_before = Pt(0)
            formatting.space_after = Pt(gaps.get(id(paragraph), 0))
            formatting.line_spacing = (pitch / metrics.line_height(size)) if pitch else 1.0
            formatting.keep_with_next = False
            formatting.keep_together = False
            formatting.widow_control = False
            properties = paragraph.get_or_add_pPr()
            number = properties.find(qn("w:numPr"))
            if number is not None:
                level = number.find(qn("w:ilvl"))
                depth = int(level.get(qn("w:val"), "0")) if level is not None else 0
                indent = size * 2 * (depth + 1)
                formatting.left_indent = Pt(indent)
                formatting.first_line_indent = Pt(-size * 2)
                formatting.tab_stops.clear_all()
                formatting.tab_stops.add_tab_stop(Pt(indent))
            # Number markers inherit the paragraph mark, not its text runs.
            mark = properties.find(qn("w:rPr"))
            if mark is None:
                mark = xml_element("w:rPr")
                properties.append(mark)
            for tag in ("w:sz", "w:szCs"):
                old = mark.find(qn(tag))
                if old is not None:
                    mark.remove(old)
                mark.append(xml_element(tag, **{qn("w:val"): str(round(size * 2))}))
            strong = [unicodedata.bidirectional(char) for char in visible_text(paragraph)
                      if unicodedata.bidirectional(char) in {"L", "R", "AL"}]
            if strong and sum(value in {"R", "AL"} for value in strong) > len(strong) / 2:
                if properties.find(qn("w:bidi")) is None:
                    properties.insert_element_before(xml_element("w:bidi"), "w:spacing", "w:ind", "w:jc", "w:rPr")
                wrapper.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def format_region(nodes: list[Any], geometry: tuple[int, int, int, int], region: Region, parent: Any,
                  warnings: list[str], metrics: FontMetrics, crop: Image.Image,
                  preferred_size: float | None = None, pitch: float | None = None) -> None:
    width, height = geometry[2] / EMU_PER_POINT, geometry[3] / EMU_PER_POINT
    if region.category in {"Title", "Section-header"}:
        # Leave a small advance allowance for heading font substitution. A
        # single wrapped word otherwise slips under the next positioned object.
        width *= 0.97
    for node in nodes:
        for table in node.iter(qn("w:tbl")):
            fit_table(table, width, crop, region, warnings, metrics)
    size = round((preferred_size or (18.0 if region.category == "Title" else 11.0)) * 2) / 2
    initial_size = size
    paragraphs = [node for node in nodes if node.tag == qn("w:p")]
    paragraph_gaps: dict[int, float] = {}
    lines = text_lines(crop)
    if len(paragraphs) > 1 and len(lines) >= len(paragraphs) and not any(
        item.find("./" + qn("w:pPr") + "/" + qn("w:numPr")) is not None for item in paragraphs
    ):
        gaps = [(index, (lines[index + 1][0] - lines[index][0]) * width / crop.width)
                for index in range(len(lines) - 1)]
        selected = sorted(sorted(gaps, key=lambda item: item[1], reverse=True)[:len(paragraphs) - 1])
        for paragraph, (_, gap) in zip(paragraphs, selected):
            paragraph_gaps[id(paragraph)] = max(0, gap - (pitch or metrics.line_height(size)))
    format_paragraphs(nodes, parent, size, metrics, paragraph_gaps, pitch)
    # A source box encloses visible glyphs, not the font's ascent/descent. The
    # caller reserves available whitespace below it before fitting.
    minimum_size = 4 if region.category == "Table" else min(6, max(4, int(size * 1.5) / 2))
    if height < metrics.line_height(6):
        minimum_size = 4
    while size > minimum_size and estimated_height(nodes, width, size, metrics, pitch) > height:
        size = max(minimum_size, size - 0.5)
        pitch = None  # Once shrinking is necessary, use measured font leading.
        format_paragraphs(nodes, parent, size, metrics, paragraph_gaps, pitch)
    if initial_size - size >= 1:
        warnings.append(f"{region.label}: reduced text from {initial_size:g} to {size:g} pt to fit; check wrapping")
    if estimated_height(nodes, width, size, metrics, pitch) > height:
        warnings.append(f"{region.label}: text may overflow its source box at {size:g} pt; content remains in DOCX but may be clipped in rendering")
    if size < 6:
        warnings.append(f"{region.label}: small {size:g} pt text required by source density; review readability")
    color = text_color(crop) if region.category != "Table" else None
    for node in nodes:
        for run_node in node.iter(qn("w:r")):
            run = Run(run_node, parent)
            run.font.name = metrics.family
            run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = RGBColor(*color)
            if any(unicodedata.bidirectional(char) in {"R", "AL"} for char in run.text):
                run.font.rtl = True
        for math_run in node.iter(qn("m:r")):
            properties = math_run.find(qn("w:rPr"))
            if properties is None:
                properties = xml_element("w:rPr")
                math_properties = math_run.find(qn("m:rPr"))
                math_run.insert(1 if math_properties is not None else 0, properties)
            for old in list(properties):
                if old.tag in {qn("w:sz"), qn("w:szCs")}:
                    properties.remove(old)
            for name in ("w:sz", "w:szCs"):
                element = xml_element(name)
                element.set(qn("w:val"), str(round(size * 2)))
                properties.append(element)


def fit_table(table: Any, width: float, crop: Image.Image, region: Region, warnings: list[str], metrics: FontMetrics) -> None:
    """Scale Pandoc's column widths; cell merging and contents remain engine-owned."""
    grid = table.findall("./" + qn("w:tblGrid") + "/" + qn("w:gridCol"))
    if not grid:
        return
    old_widths = [max(1, int(column.get(qn("w:w"), "1"))) for column in grid]
    rows = table.findall(qn("w:tr"))
    measured = table_geometry(crop, len(grid), len(rows))
    column_rules = measured.vertical.sum(axis=0).max() > crop.height * 0.18
    if measured.rows is None:
        warnings.append(f"{region.label}: table row boundaries uncertain; row heights and cell appearance could not be recovered")
    if measured.columns is not None:
        old_widths = [b - a for a, b in zip(measured.columns, measured.columns[1:])]
    elif region.category == "Table" and not any(value in region.text.lower() for value in ("<col", "width=")):
        old_widths = table_content_widths(rows, len(grid), metrics)
        warnings.append(f"{region.label}: table column boundaries uncertain; using content-proportional widths; verify column placement and OCR row associations")
    source = etree.HTML(region.text)
    if source is not None and not source.xpath('//td[@rowspan] | //th[@rowspan]'):
        counts = [sum(int(cell.get('colspan', '1')) for cell in row if cell.tag in {'td', 'th'})
                  for row in source.xpath('//tr')]
        if len(set(counts)) > 1:
            warnings.append(f"{region.label}: inconsistent OCR table row widths; review missing cells and shifted descriptions against the source (recognized cells are unchanged)")
    widths = [max(1, round(width * 20 * value / sum(old_widths))) for value in old_widths]
    for column, value in zip(grid, widths):
        column.set(qn("w:w"), str(value))
    table_width = table.find("./" + qn("w:tblPr") + "/" + qn("w:tblW"))
    if table_width is not None:
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(sum(widths)))
    properties = table.find(qn("w:tblPr"))
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = xml_element("w:tblLayout")
        # tblLayout precedes cell margins/look in the WordprocessingML schema.
        following = next((child for child in properties
                          if child.tag in {qn("w:tblCellMar"), qn("w:tblLook")}), None)
        if following is None:
            properties.append(layout)
        else:
            following.addprevious(layout)
    layout.set(qn("w:type"), "fixed")
    table_margins = properties.find(qn("w:tblCellMar"))
    if table_margins is not None:
        properties.remove(table_margins)
    dense = len(rows) > 25
    table_margins = xml_element("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margin = xml_element(f"w:{side}")
        margin.set(qn("w:w"), ("0" if dense else "10") if side in {"top", "bottom"} else ("20" if dense else "40"))
        margin.set(qn("w:type"), "dxa")
        table_margins.append(margin)
    properties.insert_element_before(table_margins, "w:tblLook", "w:tblCaption", "w:tblDescription")
    points_per_pixel = width / crop.width
    for row_index, row in enumerate(rows):
        if measured.rows is not None:
            row_properties = row.find(qn("w:trPr"))
            if row_properties is None:
                row_properties = xml_element("w:trPr")
                row.insert(0, row_properties)
            row_height = xml_element("w:trHeight")
            row_height.set(qn("w:val"), str(round((measured.rows[row_index + 1] - measured.rows[row_index]) * points_per_pixel * 20)))
            row_height.set(qn("w:hRule"), "atLeast")
            row_properties.append(row_height)
        index = 0
        for cell in row.findall(qn("w:tc")):
            properties = cell.find(qn("w:tcPr"))
            span = properties.find(qn("w:gridSpan"))
            count = int(span.get(qn("w:val"))) if span is not None else 1
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = xml_element("w:tcW")
                properties.insert(0, cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(sum(widths[index : index + count])))
            # Keep a small explicit cell inset; native tables otherwise inherit
            # margins that were not included in the OCR bounding box.
            margins = xml_element("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = xml_element(f"w:{side}")
                margin.set(qn("w:w"), ("0" if dense else "10") if side in {"top", "bottom"} else ("20" if dense else "40"))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            old_margins = properties.find(qn("w:tcMar"))
            if old_margins is not None:
                properties.remove(old_margins)
            properties.insert_element_before(margins, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers")
            if measured.columns is not None and measured.rows is not None and index + count <= len(grid):
                last_row = row_index + 1
                merge = properties.find(qn("w:vMerge"))
                if merge is not None and merge.get(qn("w:val")) == "restart":
                    for following in rows[row_index + 1:]:
                        column_index = 0
                        continuation = False
                        for other in following.findall(qn("w:tc")):
                            other_props = other.find(qn("w:tcPr"))
                            other_span = other_props.find(qn("w:gridSpan"))
                            other_count = int(other_span.get(qn("w:val"))) if other_span is not None else 1
                            if column_index == index:
                                other_merge = other_props.find(qn("w:vMerge"))
                                continuation = other_merge is not None and other_merge.get(qn("w:val"), "continue") == "continue"
                                break
                            column_index += other_count
                        if not continuation:
                            break
                        last_row += 1
                    align = properties.find(qn("w:vAlign"))
                    if align is None:
                        align = xml_element("w:vAlign")
                        properties.insert_element_before(align, "w:hideMark", "w:headers")
                    align.set(qn("w:val"), "center")
                box = (measured.columns[index], measured.rows[row_index], measured.columns[index + count], measured.rows[last_row])
                fill, rules = cell_appearance(measured, box)
                alignment = cell_alignment(measured, box)
                if alignment:
                    for paragraph in cell.findall(qn("w:p")):
                        props = paragraph.get_or_add_pPr()
                        justify = props.get_or_add_jc()
                        justify.set(qn("w:val"), alignment)
                for old in list(properties):
                    if old.tag in {qn("w:tcBorders"), qn("w:shd")}:
                        properties.remove(old)
                borders = xml_element("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    border = xml_element(f"w:{side}")
                    border.set(qn("w:val"), "single" if side in rules else "nil")
                    if side in rules:
                        color, thickness = rules[side]
                        border.set(qn("w:color"), color)
                        border.set(qn("w:sz"), str(max(2, min(96, round(thickness * points_per_pixel * 8)))))
                    borders.append(border)
                properties.insert_element_before(borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
                if fill and (column_rules or len(rules) >= 2):
                    shading = xml_element("w:shd")
                    shading.set(qn("w:fill"), fill)
                    shading.set(qn("w:val"), "clear")
                    properties.insert_element_before(shading, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
            index += count


def table_content_widths(rows: list[Any], columns: int, metrics: FontMetrics) -> list[float]:
    """Give long descriptions more room when pixel boundaries are unavailable."""
    widths = [metrics.width("000", 6) + 4 for _ in range(columns)]
    for row in rows:
        index = 0
        for cell in row.findall(qn("w:tc")):
            span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
            count = int(span.get(qn("w:val"))) if span is not None else 1
            if count == 1 and index < columns:
                lines = [visible_text(p) for p in cell.findall(qn("w:p"))]
                advance = max((metrics.width(line, 6) for line in lines), default=0)
                widths[index] = max(widths[index], min(advance, 240) + 4)
            index += count
    return widths


def take_regions(document: Any, count: int) -> list[list[Any]]:
    groups: list[list[Any]] = [[] for _ in range(count)]
    current = None
    for node in list(document._element.body):
        if node.tag == qn("w:sectPr"):
            continue
        style = node.find("./" + qn("w:pPr") + "/" + qn("w:pStyle"))
        if style is not None and style.get(qn("w:val")) == "OcrBoundary":
            current = int(visible_text(node))
        elif current is not None:
            groups[current].append(node)
        document._element.body.remove(node)
    if count and current != count - 1:
        raise ValueError("Pandoc did not preserve OCR region boundaries")
    return groups


def dotted_numbering(document: Any, nodes: list[Any], source: str) -> None:
    """Give an explicitly dotted list its own native multilevel numbering definition."""
    labels = dotted_list(source)
    if not labels:
        return
    paragraphs = [paragraph for node in nodes for paragraph in node.iter(qn("w:p"))]
    numbers = [paragraph.find("./" + qn("w:pPr") + "/" + qn("w:numPr")) for paragraph in paragraphs]
    if len(numbers) != len(labels) or any(item is None for item in numbers):
        return
    numbering = document.part.numbering_part.element
    source_id = numbers[0].find(qn("w:numId")).get(qn("w:val"))
    source_num = next(item for item in numbering.findall(qn("w:num")) if item.get(qn("w:numId")) == source_id)
    source_abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    template = next(item for item in numbering.findall(qn("w:abstractNum")) if item.get(qn("w:abstractNumId")) == source_abstract_id)
    abstract_id = max(int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))) + 1
    num_id = max(int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))) + 1
    abstract = deepcopy(template)
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = abstract.find(qn("w:nsid"))
    if nsid is not None:
        nsid.set(qn("w:val"), f"{abstract_id:08X}")
    for level in abstract.findall(qn("w:lvl")):
        depth = int(level.get(qn("w:ilvl")))
        level.find(qn("w:numFmt")).set(qn("w:val"), "decimal")
        level.find(qn("w:lvlText")).set(qn("w:val"), ".".join(f"%{index + 1}" for index in range(depth + 1)) + ".")
        start = level.find(qn("w:start"))
        if start is None:
            start = xml_element("w:start")
            level.insert(0, start)
        start.set(qn("w:val"), labels[0][1] if depth == 0 else "1")
    numbering.insert_element_before(abstract, "w:num", "w:numIdMacAtCleanup")
    number = xml_element("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = xml_element("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    numbering.insert_element_before(number, "w:numIdMacAtCleanup")
    for properties in numbers:
        properties.find(qn("w:numId")).set(qn("w:val"), str(num_id))


def text_geometry(page: Page, region: Region, size: float | None) -> tuple[int, int, int, int]:
    """Reserve font descent/leading in free space without crossing another region."""
    left, top, width, height = region_geometry(page, region)
    if size is None:
        return left, top, width, height
    x1, _, x2, y2 = region.bbox
    next_top = min((other.bbox[1] for other in page.regions
                    if other is not region and other.bbox[1] >= y2
                    and other.bbox[0] < x2 and other.bbox[2] > x1), default=page.height)
    scale = page_geometry(page)[2]
    reserve = min(size * 0.65 * EMU_PER_POINT, max(0, (next_top - y2 - 2) * scale))
    return left, top, width, height + round(reserve)


def formula_geometry(page: Page, region: Region, geometry: tuple[int, int, int, int]) -> tuple[int, int, int, int]:
    """Widen a centered display equation into free space beside it.

    Word breaks an equation that is wider than its text box onto a new line,
    which the box then clips. Font substitution makes that width hard to predict.
    """
    x1, y1, x2, y2 = region.bbox
    beside = [other.bbox for other in page.regions if other is not region and other.bbox[1] < y2 and other.bbox[3] > y1]
    left = max((bbox[2] for bbox in beside if bbox[2] <= x1), default=0)
    right = min((bbox[0] for bbox in beside if bbox[0] >= x2), default=page.width)
    margin = round(min(x1 - left, right - x2) * page_geometry(page)[2])
    return geometry[0] - margin, geometry[1], geometry[2] + 2 * margin, geometry[3]


def avoid_graphics(page: Page, boxes: list[tuple[int, int, int, int]], warnings: list[str]) -> tuple[Page, list[tuple[int, int, int, int]]]:
    """Reserve space for edge graphics; do not cover editable text with an uncertain crop."""
    regions = list(page.regions)
    accepted = []
    for box in boxes:
        gx1, gy1, gx2, gy2 = box
        adjusted = list(regions)
        blocked = False
        for index, region in enumerate(regions):
            x1, y1, x2, y2 = region.bbox
            if region.category == "Picture" or gx2 <= x1 or gx1 >= x2 or gy2 <= y1 or gy1 >= y2:
                continue
            if region.category in {"Text", "Caption", "Footnote"} and gx2 <= x1 + (x2 - x1) * 0.25:
                adjusted[index] = replace(region, bbox=(gx2 + 3, y1, x2, y2))
            elif region.category in {"Text", "Caption", "Footnote"} and gx1 >= x2 - (x2 - x1) * 0.25:
                adjusted[index] = replace(region, bbox=(x1, y1, gx1 - 3, y2))
            else:
                blocked = True
        if blocked:
            warnings.append(f"{page.regions[0].label.split(',')[0] if page.regions else page.path.name}: graphic candidate at {box} overlaps editable content; crop omitted for review")
            continue
        for before, after in zip(regions, adjusted):
            if before.bbox != after.bbox:
                warnings.append(f"{before.label}: narrowed text box around recovered graphic at {box}; check wrapping")
        regions = adjusted
        accepted.append(box)
    return replace(page, regions=regions), accepted


def orient_regions(page: Page, image: Image.Image, metrics: FontMetrics, warnings: list[str]) -> Page:
    regions = list(page.regions)
    for index, region in enumerate(regions):
        if region.rotation or region.category not in {"Text", "Caption", "List-item"}:
            continue
        with image.crop(region.bbox) as crop:
            angle = text_rotation(crop, region.text, metrics)
        if angle:
            regions[index] = replace(region, rotation=angle)
            warnings.append(f"{region.label}: inferred {angle} degree text rotation from glyphs; verify direction")
    # A matching adjacent line can resolve a direction ambiguous in isolation.
    confirmed = [region for region in regions if region.rotation in {90, 270}]
    for index, region in enumerate(regions):
        x1, y1, x2, y2 = region.bbox
        if region.rotation or region.category != "Text" or y2 - y1 < (x2 - x1) * 3:
            continue
        neighbors = [other.rotation for other in confirmed
                     if abs(other.bbox[0] - x1) < (x2 - x1) * 2
                     and min(other.bbox[3], y2) - max(other.bbox[1], y1) > min(y2 - y1, other.bbox[3] - other.bbox[1]) * 0.8]
        if neighbors and len(set(neighbors)) == 1:
            regions[index] = replace(region, rotation=neighbors[0])
            warnings.append(f"{region.label}: inferred {neighbors[0]} degree rotation from adjacent text; verify direction")
        else:
            warnings.append(f"{region.label}: narrow tall text box; orientation is ambiguous, retained horizontal text")
    return replace(page, regions=regions)


def export_word(pages: Sequence[Page], output: Path, *, font: str = "Arial", recover_graphics: bool = True,
                deskew: bool = True) -> list[str]:
    if not pages:
        raise ValueError("At least one page is required")
    if output.resolve() in {page.path.resolve() for page in pages}:
        raise ValueError("Output must not overwrite a source image")
    warnings: list[str] = []
    metrics = FontMetrics(font)
    if metrics.resolved_family.casefold() != font.casefold():
        warnings.append(f"Document: measuring {font} with installed {metrics.resolved_family}; font substitution can change wrapping")
    regions = [region for page in pages for region in page.regions]
    with tempfile.TemporaryDirectory(prefix="ocr-word-") as temporary:
        content_path = convert_regions(regions, Path(temporary), warnings)
        document = Document(content_path)
        groups = take_regions(document, len(regions))
        normal = document.styles["Normal"]
        normal.font.name = font
        normal.font.size = Pt(11)
        region_index = 0
        identifier = 0
        for page_index, page in enumerate(pages):
            section = document.sections[0] if page_index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
            width, height, _, _, _ = page_geometry(page)
            section.orientation = WD_ORIENT.LANDSCAPE if page.width > page.height else WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = width, height
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(0)
            section.header_distance = section.footer_distance = Mm(0)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            with prepared_page(page, warnings, deskew=deskew) as (page, image):
                page = orient_regions(page, image, metrics, warnings)
                boxes, residual = missing_graphics(image, page.regions)
                if recover_graphics:
                    page, boxes = avoid_graphics(page, boxes, warnings)
                estimates = []
                points_per_pixel = page_geometry(page)[2] / EMU_PER_POINT
                for region in page.regions:
                    with image.crop(region.bbox) as crop:
                        if region.rotation:
                            crop = crop.rotate(region.rotation, expand=True)
                        estimates.append(source_font_size(crop, points_per_pixel, metrics,
                                                          region.text if region.category != "Table" else None)
                                         if region.category not in {"Picture", "Formula"} else (None, None))
                body_sizes = [size for region, (size, _) in zip(page.regions, estimates)
                              if region.category == "Text" and size is not None]
                body_size = median(body_sizes) if body_sizes else 11.0
                for region, (size, pitch) in zip(page.regions, estimates):
                    identifier += 1
                    geometry = region_geometry(page, region)
                    if region.category == "Picture":
                        add_picture(paragraph, identifier, region, geometry, image, page.regions, warnings)
                        if geometry[2] * geometry[3] > width * height * 0.25:
                            warnings.append(f"{region.label}: large Picture region; any text or form fields inside remain pixels, not editable content")
                    else:
                        nodes = groups[region_index]
                        dotted_numbering(document, nodes, region.text)
                        if region.category == "Text" and size is not None and pitch is not None and abs(size - body_size) <= 1:
                            size = body_size
                        if region.category == "Table" and size is not None and size > body_size * 1.4:
                            warnings.append(f"{region.label}: table font measurement may include rules or fills; using page body size {body_size:g} pt")
                            size, pitch = body_size, None
                        if region.category == "Formula":
                            size = body_size
                        geometry = region_geometry(page, region) if region.rotation else text_geometry(page, region, size)
                        fitting = (*geometry[:2], geometry[3], geometry[2]) if region.rotation in {90, 270} else geometry
                        with image.crop(region.bbox) as crop:
                            if region.rotation:
                                crop = crop.rotate(region.rotation, expand=True)
                            format_region(nodes, fitting, region, document._body, warnings, metrics, crop, size, pitch)
                        if region.category == "Formula" and not region.rotation:
                            geometry = formula_geometry(page, region, geometry)
                        add_textbox(paragraph, identifier, region, geometry, nodes)
                    region_index += 1
                for index, box in enumerate(boxes, start=1):
                    label = f"Page {page_index + 1}, recovered graphic {index}"
                    if recover_graphics:
                        identifier += 1
                        region = Region(box, "Picture", "", label)
                        add_picture(paragraph, identifier, region, region_geometry(page, region), image)
                    action = "cropped" if recover_graphics else "detected (recovery disabled)"
                    warnings.append(f"{label}: {action} at {box}; verify graphic classification and text overlap")
                if residual > max(50, page.width * page.height * 0.0001):
                    warnings.append(f"Page {page_index + 1}: {residual} ink pixels outside OCR regions remain unrepresented; review missing text/graphics")
        output.parent.mkdir(parents=True, exist_ok=True)
        descriptor, name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".docx", dir=output.parent)
        os.close(descriptor)
        try:
            document.save(name)
            os.replace(name, output)
        finally:
            Path(name).unlink(missing_ok=True)
    return warnings
