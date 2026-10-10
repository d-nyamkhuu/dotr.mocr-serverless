"""Exercise the real Pandoc engine and the OCR-to-Word placement adapter."""

from __future__ import annotations

import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.shared import Mm
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

from word_export.export import (
    estimated_height,
    export_word,
    load_pages,
    region_geometry,
)
from word_export.content import normalize_math, normalize_table
from word_export.layout import FontMetrics, missing_graphics, source_font_size

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
}


def element(category, text=None, bbox=(30, 40, 550, 160)):
    return {"category": category, "text": text, "bbox": list(bbox)}


def response(width=600, height=800, elements=None):
    return {
        "status": "ok",
        "image": {"original_width": width, "original_height": height},
        "elements": elements if elements is not None else [element("Text", "Sample")],
    }


class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.image = self.root / "portrait.png"
        with Image.new("RGB", (600, 800), "white") as image:
            image.save(self.image)
        self.output = self.root / "output.docx"

    def convert(self, payload, paths=None):
        pages = load_pages(payload, paths or [self.image])
        warnings = export_word(pages, self.output)
        with ZipFile(self.output) as archive:
            xml = etree.fromstring(archive.read("word/document.xml"))
        return xml, warnings

    def test_native_content_and_equations(self):
        with Image.open(self.image) as image:
            image.paste((220, 20, 40), (30, 400, 130, 500))
            image.save(self.image)
        table = (
            '<table><tr><th rowspan="2">Merged header</th><th colspan="2">Group</th></tr>'
            '<tr><td>Өгөгдөл</td><td>42</td></tr></table>'
        )
        payload = response(elements=[
            element("Title", "Монгол гарчиг", (20, 10, 580, 60)),
            element("Text", "Latin **bold** and *italic*.\nКирилл Өө Үү $x^2$.", (20, 70, 580, 170)),
            element("List-item", "- First\n- Second", (20, 180, 580, 250)),
            element("Table", table, (20, 260, 580, 380)),
            element("Picture", bbox=(30, 400, 130, 500)),
            element("Formula", r"\frac{a}{b}+\sqrt{x}+x_1^2", (20, 510, 580, 590)),
            element("Formula", r"\begin{pmatrix}a&b\\c&d\end{pmatrix}", (20, 600, 580, 700)),
        ])
        xml, warnings = self.convert(payload)
        self.assertTrue(all("font substitution" in warning or "table column boundaries uncertain" in warning or "table row boundaries uncertain" in warning for warning in warnings), warnings)
        text = "".join(xml.xpath("//w:t/text()", namespaces=NS))
        for value in ["Монгол гарчиг", "Кирилл Өө Үү", "Merged header", "Өгөгдөл", "First", "Second"]:
            self.assertIn(value, text)
        for path in ["//w:b", "//w:i", "//w:numPr", "//w:br", "//m:f", "//m:rad", "//m:sSubSup", "//m:m"]:
            self.assertTrue(xml.xpath(path, namespaces=NS), path)
        self.assertEqual(len(xml.xpath("//w:tbl", namespaces=NS)), 1)
        self.assertTrue(xml.xpath('//w:gridSpan[@w:val="2"]', namespaces=NS))
        self.assertTrue(xml.xpath('//w:vMerge[@w:val="restart"]', namespaces=NS))
        self.assertTrue(xml.xpath('//w:vMerge[@w:val="continue"]', namespaces=NS))
        with ZipFile(self.output) as archive:
            pictures = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(pictures), 1)
            with Image.open(io.BytesIO(archive.read(pictures[0]))) as crop:
                self.assertEqual(crop.size, (100, 100))
                self.assertEqual(crop.getpixel((50, 50)), (220, 20, 40))
        # python-docx can reopen the final package with all styles/relationships intact.
        self.assertEqual(len(Document(self.output).sections), 1)

    def test_multiple_pages_order_orientation_and_coordinates(self):
        landscape = self.root / "landscape.png"
        Image.new("RGB", (800, 600), "white").save(landscape)
        payload = [
            response(elements=[element("Page-header", "First page")]),
            response(800, 600, [element("Page-header", "Second page")]),
        ]
        pages = load_pages(payload, [self.image, landscape])
        xml, _ = self.convert(payload, [self.image, landscape])
        self.assertEqual(xml.xpath("//w:t/text()", namespaces=NS), ["First page", "Second page"])
        sections = Document(self.output).sections
        self.assertEqual(len(sections), 2)
        # Word dimensions are serialized in twips, so allow one twip rounding.
        for section, width, height in [(sections[0], 210, 297), (sections[1], 297, 210)]:
            self.assertLess(abs(section.page_width - Mm(width)), 635)
            self.assertLess(abs(section.page_height - Mm(height)), 635)
        anchors = xml.xpath("//wp:anchor", namespaces=NS)
        self.assertEqual(len(anchors), 2)
        self.assertEqual(len({anchor.find("wp:docPr", NS).get("id") for anchor in anchors}), 2)
        for page, anchor in zip(pages, anchors):
            x, y, width, height = region_geometry(page, page.regions[0])
            self.assertEqual(anchor.find("wp:positionH/wp:posOffset", NS).text, str(x))
            self.assertEqual(anchor.find("wp:positionV/wp:posOffset", NS).text, str(y))
            self.assertEqual(anchor.find("wp:extent", NS).get("cx"), str(width))
            self.assertEqual(anchor.find("wp:extent", NS).get("cy"), str(height))
        self.assertGreater(region_geometry(pages[0], pages[0].regions[0])[1], Mm(40 * 210 / 600))

    def test_source_fallbacks_and_overflow_are_reported(self):
        formula = r"\notARealCommand{x}"
        table = "<table><tr><td>truncated"
        long_text = "Өгүүлбэр " * 100
        xml, warnings = self.convert(response(elements=[
            element("Formula", formula),
            element("Table", table, (30, 170, 550, 300)),
            element("Text", long_text, (30, 320, 35, 325)),
        ]))
        text = "".join(xml.xpath("//w:t/text()", namespaces=NS))
        self.assertIn(formula, text)
        self.assertIn(table, text)
        self.assertEqual(text.count("Өгүүлбэр"), 100)
        self.assertTrue(any("Page 1, region 1" in warning and "LaTeX" in warning for warning in warnings))
        self.assertTrue(any("Page 1, region 2" in warning and "HTML" in warning for warning in warnings))
        self.assertTrue(any("Page 1, region 3" in warning and "overflow" in warning for warning in warnings))
        self.assertFalse(xml.xpath("//w:tbl", namespaces=NS))

    def test_inline_failure_and_backslash_math(self):
        xml, warnings = self.convert(response(elements=[
            element("Text", r"Equation \(x_1\), then $\notARealCommand{x}$.")
        ]))
        self.assertTrue(xml.xpath("//m:oMath", namespaces=NS))
        self.assertTrue(any("Page 1, region 1" in warning for warning in warnings))
        self.assertIn(r"\notARealCommand{x}", "".join(xml.xpath("//w:t/text()", namespaces=NS)))

    def test_table_math_and_explicit_column_widths(self):
        xml, warnings = self.convert(response(elements=[element("Table", (
            '<table><colgroup><col width="25%"><col width="75%"></colgroup>'
            '<tr><th>Variable</th><th>Equation</th></tr>'
            '<tr><td>$x_1$</td><td>$\\frac{a}{b}$</td></tr></table>'
        ))]))
        self.assertTrue(all("font substitution" in warning or "table column boundaries uncertain" in warning or "table row boundaries uncertain" in warning for warning in warnings), warnings)
        self.assertTrue(xml.xpath("//w:tc//m:f", namespaces=NS))
        self.assertTrue(xml.xpath('//w:tblLayout[@w:type="fixed"]', namespaces=NS))
        widths = [int(value) for value in xml.xpath("//w:gridCol/@w:w", namespaces=NS)]
        self.assertAlmostEqual(widths[1] / widths[0], 3, places=2)

    def test_untrusted_markup_cannot_embed_external_images_or_xml(self):
        xml, warnings = self.convert(response(elements=[
            element("Text", '![Alt text](https://example.invalid/image.png)\n\n`<w:p/>`{=openxml}')
        ]))
        self.assertIn("Alt text", "".join(xml.xpath("//w:t/text()", namespaces=NS)))
        self.assertFalse(xml.xpath("//a:blip", namespaces=NS))
        self.assertTrue(any("inline image" in warning for warning in warnings))

    def test_exif_orientation_is_applied_before_cropping(self):
        source = self.root / "rotated.jpg"
        exif = Image.Exif()
        exif[274] = 6
        Image.new("RGB", (80, 40), "blue").save(source, exif=exif)
        payload = response(40, 80, [element("Picture", bbox=(0, 0, 40, 80))])
        self.convert(payload, [source])
        with ZipFile(self.output) as archive:
            name = next(name for name in archive.namelist() if name.startswith("word/media/"))
            with Image.open(io.BytesIO(archive.read(name))) as image:
                self.assertEqual(image.size, (40, 80))

    def test_result_formats(self):
        result = response()
        for payload in (result, [result], [{"index": 0, "response": result, "error": None}]):
            with self.subTest(payload=payload):
                self.assertEqual(load_pages(payload, [self.image])[0].regions[0].text, "Sample")

    def test_invalid_inputs(self):
        invalid = [None, [], [response(), response()], {"status": "degraded", "elements": None},
                   [{"response": None, "error": "failed"}], [{"index": 2, "response": response()}]]
        mismatch = response(500, 800)
        invalid.append(mismatch)
        for bbox in ([10, 0, 5, 10], [0, 0, 601, 10], [-1, 0, 5, 10], [0, 0, 0, 10], [0, 0, 5.5, 10]):
            invalid.append(response(elements=[element("Text", "bad", bbox)]))
        invalid.append(response(elements=[element("Text", "bad\x00text")]))
        for payload in invalid:
            with self.subTest(payload=payload), self.assertRaises((TypeError, ValueError)):
                load_pages(payload, [self.image])
        self.assertFalse(self.output.exists())

    def test_empty_page_is_preserved(self):
        self.convert([response(elements=[]), response()], [self.image, self.image])
        self.assertEqual(len(Document(self.output).sections), 2)

    def test_malformed_spans_fall_back(self):
        for source in (
            '<table><tr><td colspan="no">bad</td></tr></table>',
            '<table><tr><td rowspan="3">bad</td></tr></table>',
            '<table><tr><td><table><tr><td>nested</td></tr></table></td></tr></table>',
        ):
            with self.subTest(source=source), self.assertRaises(ValueError):
                normalize_table(source)

    def test_atomic_output_preserves_existing_file_on_failure(self):
        self.output.write_bytes(b"existing document")
        pages = load_pages(response(), [self.image])
        with patch("docx.document.Document.save", side_effect=OSError("disk failure")), self.assertRaises(OSError):
            export_word(pages, self.output)
        self.assertEqual(self.output.read_bytes(), b"existing document")
        self.assertEqual(list(self.root.glob(".output.docx.*")), [])

    def test_safe_html_styles_and_unbalanced_markup(self):
        xml, warnings = self.convert(response(elements=[element("Text", (
            'H<sub>2</sub>O x<sup><b>2</b></sup> <u>under</u> <s>gone</s> '
            '<sup>unclosed <script>literal()</script> <sup onclick="bad">unsafe</sup>'
        ))]))
        self.assertTrue(xml.xpath('//w:vertAlign[@w:val="superscript"]', namespaces=NS))
        self.assertTrue(xml.xpath('//w:vertAlign[@w:val="subscript"]', namespaces=NS))
        self.assertTrue(xml.xpath('//w:u', namespaces=NS))
        self.assertTrue(xml.xpath('//w:strike', namespaces=NS))
        text = ''.join(xml.xpath('//w:t/text()', namespaces=NS))
        self.assertIn('<sup>unclosed', text)
        self.assertIn('<script>literal()</script>', text)
        self.assertIn('onclick="bad"', text)
        self.assertTrue(any('unsupported markup' in value for value in warnings))

    def test_scripted_parentheses_are_editable_delimiters(self):
        source = r'\hat{v}^i = (R^{-1})_j^i v^j'
        normalized = normalize_math(source)
        self.assertIn(r'\left(R^{-1}\right)_j^i', normalized)
        self.assertEqual(normalize_math(normalized), normalized)
        for literal in [r'\text{(hello)_a}', r'\left(x\right)^2', r'\(x\)', r'(a+b)', r'\bigl(x\bigr)^2']:
            self.assertEqual(normalize_math(literal), literal)
        xml, _ = self.convert(response(elements=[element('Formula', source)]))
        self.assertTrue(xml.xpath('//m:d', namespaces=NS))
        self.assertIn('R', ''.join(xml.xpath('//m:t/text()', namespaces=NS)))

    def test_font_metrics_and_source_line_size(self):
        metrics = FontMetrics()
        self.assertGreater(metrics.width('WWWӨӨ', 12), metrics.width('iiiii', 12) * 2)
        image = Image.new('RGB', (800, 150), 'white')
        font = ImageFont.truetype(metrics.paths[False, False], 28)
        draw = ImageDraw.Draw(image)
        draw.text((5, 2), 'Latin Кирилл Өө Үү', font=font, fill='black')
        draw.text((5, 42), 'Second line of text', font=font, fill='black')
        size, pitch = source_font_size(image, 0.5, metrics)
        self.assertAlmostEqual(size, 14, delta=2)
        self.assertAlmostEqual(pitch, 20, delta=2)
        small = Document().add_paragraph('iiiiiiii iiiiiii iiiii')._p
        wide = Document().add_paragraph('WWWWWWWW WWWWWWW WWWWW')._p
        self.assertGreater(estimated_height([wide], 80, 12, metrics), estimated_height([small], 80, 12, metrics))

    def test_detected_table_dimensions_rules_and_fill(self):
        with Image.open(self.image) as image:
            draw = ImageDraw.Draw(image)
            draw.rectangle((30, 40, 549, 99), fill='#D0E0F0')
            for x in (30, 160, 549):
                draw.line((x, 40, x, 239), fill='black', width=2)
            for y in (40, 100, 239):
                draw.line((30, y, 549, y), fill='black', width=2)
            image.save(self.image)
        xml, _ = self.convert(response(elements=[element('Table', '<table><tr><th>A</th><th>B</th></tr><tr><td>C</td><td>D</td></tr></table>', (30, 40, 550, 240))]))
        widths = [int(x) for x in xml.xpath('//w:gridCol/@w:w', namespaces=NS)]
        self.assertAlmostEqual(widths[1] / widths[0], 3, delta=0.06)
        heights = [int(x) for x in xml.xpath('//w:trHeight/@w:val', namespaces=NS)]
        self.assertAlmostEqual(heights[1] / heights[0], 140 / 60, delta=0.1)
        self.assertTrue(xml.xpath('//w:trHeight[@w:hRule="atLeast"]', namespaces=NS))
        self.assertTrue(xml.xpath('//w:tcBorders/w:top[@w:val="single"]', namespaces=NS))
        self.assertTrue(xml.xpath('//w:shd[@w:fill="D0E0F0"]', namespaces=NS))

    def test_graphics_recovery_and_ambiguous_ink(self):
        with Image.open(self.image) as image:
            draw = ImageDraw.Draw(image)
            draw.polygon([(40, 80), (95, 40), (95, 65), (140, 65), (140, 95), (95, 95), (95, 120)], fill='blue')
            draw.ellipse((220, 80, 250, 110), fill='green')
            # Small disjoint marks must not become image overlays.
            for x in range(30, 90, 8):
                draw.rectangle((x, 550, x + 2, 558), fill='black')
            image.save(self.image)
        payload = response(elements=[element('Text', 'Recognized text in a broad box', (30, 30, 500, 170))])
        xml, warnings = self.convert(payload)
        self.assertEqual(len(xml.xpath('//a:blip', namespaces=NS)), 1)
        self.assertTrue(any('crop omitted' in value for value in warnings))
        self.assertTrue(any('narrowed text box' in value for value in warnings))
        self.assertTrue(any('recovered graphic' in value for value in warnings))
        self.assertTrue(any('unrepresented' in value for value in warnings))
        pages = load_pages(payload, [self.image])
        export_word(pages, self.output, recover_graphics=False)
        with ZipFile(self.output) as archive:
            self.assertFalse(any(name.startswith('word/media/') for name in archive.namelist()))
        with Image.open(self.image) as image:
            covered = load_pages(response(elements=[element('Picture', bbox=(0, 0, 600, 800))]), [self.image])
            self.assertEqual(missing_graphics(image, covered[0].regions), ([], 0))

    def test_nested_list_indentation(self):
        xml, _ = self.convert(response(elements=[element('Text', '- First\n    - Nested\n        - Deeper\n- Last', (20, 20, 580, 400))]))
        levels = [int(value) for value in xml.xpath('//w:numPr/w:ilvl/@w:val', namespaces=NS)]
        self.assertGreater(max(levels), min(levels))
        indents = [int(value) for value in xml.xpath('//w:pPr/w:ind/@w:left', namespaces=NS)]
        self.assertGreater(max(indents), min(indents))

    def test_dotted_lists_keep_multilevel_native_labels(self):
        xml, _ = self.convert(response(elements=[element('Text', '1. One\n1.1. Two\n1.1.1. Three\n1.1.2. Four\n2. Five', (20, 20, 580, 500))]))
        identifiers = xml.xpath('//w:numId/@w:val', namespaces=NS)
        self.assertEqual(len(set(identifiers)), 1)
        with ZipFile(self.output) as archive:
            numbering = etree.fromstring(archive.read('word/numbering.xml'))
        number = numbering.xpath(f'//w:num[@w:numId="{identifiers[0]}"]', namespaces=NS)[0]
        abstract_id = number.find('w:abstractNumId', NS).get(f'{{{NS["w"]}}}val')
        abstract = numbering.xpath(f'//w:abstractNum[@w:abstractNumId="{abstract_id}"]', namespaces=NS)[0]
        self.assertTrue(abstract.xpath('.//w:lvlText[@w:val="%1.%2.%3."]', namespaces=NS))

    def test_display_equation_box_widens_around_its_center(self):
        # Word breaks a display equation wider than its box onto a clipped second line.
        source = r'\hat{e}_i = \sum_{j=1}^{n} e_j R_i^j = e_j R_i^j.'
        payload = response(elements=[
            element('Text', 'Side note', (20, 300, 180, 360)),
            element('Formula', source, (250, 300, 350, 360)),
        ])
        xml, _ = self.convert(payload)
        pages = load_pages(payload, [self.image])
        left, _, width, height = region_geometry(pages[0], pages[0].regions[1])
        anchor = xml.xpath('//wp:anchor[wp:docPr/@name="Page 1, region 2"]', namespaces=NS)[0]
        offset = int(anchor.xpath('./wp:positionH/wp:posOffset/text()', namespaces=NS)[0])
        extent = anchor.find('wp:extent', NS)
        self.assertEqual(offset + int(extent.get('cx')) / 2, left + width / 2)
        note_left, _, note_width, _ = region_geometry(pages[0], pages[0].regions[0])
        self.assertAlmostEqual(offset, note_left + note_width, delta=1)
        self.assertGreater(int(extent.get('cx')), width * 2)
        self.assertGreater(int(extent.get('cy')), height)

    def test_failed_normalized_math_preserves_original(self):
        source = r'(R^{-1})_j^i + \notARealCommand{x}'
        xml, warnings = self.convert(response(elements=[element('Formula', source)]))
        self.assertIn(source, ''.join(xml.xpath('//w:t/text()', namespaces=NS)))
        self.assertTrue(any('unsupported equation' in warning for warning in warnings))

    def test_colored_letters_are_not_recovered_as_pictures(self):
        metrics = FontMetrics()
        with Image.open(self.image) as image:
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype(metrics.paths[False, False], 20)
            draw.text((40, 40), 'WWW MMM Colored heading and body', font=font, fill='blue')
            image.save(self.image)
        xml, _ = self.convert(response(elements=[element('Text', 'WWW MMM Colored heading and body', (30, 30, 550, 160))]))
        self.assertFalse(xml.xpath('//a:blip', namespaces=NS))

    def test_paragraph_gaps_are_measured(self):
        metrics = FontMetrics()
        with Image.open(self.image) as image:
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype(metrics.paths[False, False], 20)
            draw.text((40, 45), 'First paragraph', font=font, fill='black')
            draw.text((40, 110), 'Second paragraph', font=font, fill='black')
            image.save(self.image)
        xml, _ = self.convert(response(elements=[element('Text', 'First paragraph\n\nSecond paragraph', (30, 40, 550, 160))]))
        spacings = [int(value) for value in xml.xpath('//w:txbxContent/w:p/w:pPr/w:spacing/@w:after', namespaces=NS)]
        self.assertGreater(spacings[0], 0)

    def test_list_fitting_uses_tabs_and_marker_font(self):
        xml, _ = self.convert(response(elements=[
            element('List-item', '12. Carrying on at twelve', (30, 40, 290, 66)),
            element('List-item', '◆ An instruction', (30, 80, 550, 120)),
        ]))
        paragraphs = xml.xpath('//w:txbxContent/w:p[w:pPr/w:numPr]', namespaces=NS)
        self.assertEqual(len(paragraphs), 2)
        self.assertNotIn('◆', ''.join(xml.xpath('//w:t/text()', namespaces=NS)))
        for paragraph in paragraphs:
            left = paragraph.xpath('./w:pPr/w:ind/@w:left', namespaces=NS)[0]
            self.assertEqual(paragraph.xpath('./w:pPr/w:tabs/w:tab/@w:pos', namespaces=NS), [left])
            size = paragraph.xpath('./w:pPr/w:rPr/w:sz/@w:val', namespaces=NS)[0]
            self.assertEqual(paragraph.xpath('./w:r/w:rPr/w:sz/@w:val', namespaces=NS), [size])

    def test_light_merged_header_geometry(self):
        from word_export.layout import table_geometry
        image = Image.new('RGB', (500, 206), 'white')
        draw = ImageDraw.Draw(image)
        draw.rectangle((5, 5, 494, 75), fill='#A0A0A0')
        for x in (5, 205, 345, 494):
            draw.line((x, 5, x, 200), fill='#D0D0D0', width=1)
        for y in (5, 75, 115, 155, 200):
            draw.line((5, y, 494, y), fill='#D0D0D0', width=1)
        draw.line((205, 50, 494, 50), fill='white', width=1)
        geometry = table_geometry(image, 3, 5)
        self.assertIsNotNone(geometry.columns)
        self.assertIsNotNone(geometry.rows)
        self.assertAlmostEqual(geometry.columns[1], 205, delta=3)
        self.assertAlmostEqual(geometry.rows[1], 50, delta=3)

    def test_rotated_text_inference_and_explicit_direction(self):
        from word_export.layout import text_rotation
        metrics = FontMetrics()
        font = ImageFont.truetype(metrics.paths[False, False], 26)
        line = Image.new('RGB', (250, 45), 'white')
        ImageDraw.Draw(line).text((5, 5), 'First paragraph.', font=font, fill='black')
        vertical = line.rotate(90, expand=True)
        self.assertEqual(text_rotation(vertical, 'First paragraph.', metrics), 270)
        with Image.open(self.image) as image:
            image.paste(vertical, (100, 100))
            image.save(self.image)
        xml, warnings = self.convert(response(elements=[element('Text', 'First paragraph.', (100, 100, 145, 350))]))
        self.assertTrue(xml.xpath('//wps:bodyPr[@vert="vert270"]', namespaces=NS))
        self.assertTrue(any('inferred 270' in warning for warning in warnings))
        item = element('Text', 'Explicit direction')
        item['rotation'] = 90
        xml, _ = self.convert(response(elements=[item]))
        self.assertTrue(xml.xpath('//wps:bodyPr[@vert="vert"]', namespaces=NS))
        for invalid in [45, True, '90']:
            item['rotation'] = invalid
            with self.assertRaisesRegex(ValueError, 'rotation'):
                load_pages(response(elements=[item]), [self.image])

    def test_combined_table_keeps_header_column_fallback(self):
        from word_export.layout import table_geometry
        metrics = FontMetrics()
        image = Image.new('RGB', (500, 300), 'white')
        draw = ImageDraw.Draw(image)
        for y in (0, 50, 100, 150, 200, 299):
            draw.line((0, y, 499, y), fill='black', width=2)
        for x in (0, 250, 499):
            draw.line((x, 50, x, 299), fill='black', width=2)
        font = ImageFont.truetype(metrics.paths[False, False], 18)
        for x, text in zip((5, 130, 255, 380), ('Name', 'Game', 'Fame', 'Blame')):
            draw.text((x, 10), text, font=font, fill='black')
        # Lower merged rows cross otherwise empty header column gutters.
        draw.text((5, 160), 'A long cell spanning more than one column', font=font, fill='black')
        geometry = table_geometry(image, 4, 5)
        self.assertIsNotNone(geometry.columns)
        self.assertAlmostEqual(geometry.columns[1], 125, delta=5)

    def test_white_source_letters_remain_visible_without_background(self):
        from word_export.layout import text_color
        metrics = FontMetrics()
        image = Image.new('RGB', (240, 50), 'red')
        ImageDraw.Draw(image).text((5, 5), 'Header date', font=ImageFont.truetype(metrics.paths[False, False], 24), fill='white')
        self.assertEqual(text_color(image), (0, 0, 0))

    def test_picture_text_overlap_replaces_only_flat_background(self):
        metrics = FontMetrics()
        with Image.open(self.image) as image:
            draw = ImageDraw.Draw(image)
            draw.rectangle((50, 50, 70, 70), fill='red')
            draw.text((100, 150), 'Editable', font=ImageFont.truetype(metrics.paths[False, False], 20), fill='blue')
            image.save(self.image)
        original = self.image.read_bytes()
        xml, warnings = self.convert(response(elements=[
            element('Picture', bbox=(20, 20, 580, 600)),
            element('Text', 'Editable', (98, 145, 195, 180)),
        ]))
        self.assertIn('Editable', ''.join(xml.xpath('//w:t/text()', namespaces=NS)))
        self.assertTrue(any('replaced duplicate' in warning for warning in warnings))
        self.assertTrue(any('remain pixels' in warning for warning in warnings))
        self.assertEqual(self.image.read_bytes(), original)
        picture_anchor = xml.xpath('//wp:anchor[.//a:blip]', namespaces=NS)[0]
        self.assertEqual(picture_anchor.get('behindDoc'), '1')
        with ZipFile(self.output) as archive:
            name = next(name for name in archive.namelist() if name.startswith('word/media/'))
            with Image.open(io.BytesIO(archive.read(name))) as crop:
                self.assertEqual(crop.getpixel((40, 40)), (255, 0, 0))
                self.assertEqual(crop.crop((78, 125, 175, 160)).getextrema(), ((255, 255),) * 3)
        _, warnings = self.convert(response(elements=[
            element('Text', 'Editable', (98, 145, 195, 180)),
            element('Picture', bbox=(150, 140, 580, 600)),
        ]))
        self.assertTrue(any('replaced duplicate' in warning for warning in warnings))
        with ZipFile(self.output) as archive:
            name = next(name for name in archive.namelist() if name.startswith('word/media/'))
            with Image.open(io.BytesIO(archive.read(name))) as crop:
                self.assertEqual(crop.crop((0, 5, 45, 40)).getextrema(), ((255, 255),) * 3)
        with Image.open(self.image) as image:
            draw = ImageDraw.Draw(image)
            for x in range(98, 195, 4):
                draw.line((x, 140, x, 185), fill='black', width=2)
            image.save(self.image)
        _, warnings = self.convert(response(elements=[
            element('Picture', bbox=(20, 20, 580, 600)),
            element('Text', 'Editable', (98, 145, 195, 180)),
        ]))
        self.assertTrue(any('unresolved picture/text overlap' in warning for warning in warnings))

    def test_multiple_roots_warn_without_changing_recognition(self):
        xml, warnings = self.convert(response(elements=[element('Formula', r'\sqrt{4\sqrt[3]{x+1}}')]))
        self.assertTrue(xml.xpath('//m:rad/m:e/m:rad', namespaces=NS))
        self.assertTrue(any('verify radical scope' in warning for warning in warnings))

    def test_rtl_paragraph_and_source_text_color(self):
        metrics = FontMetrics()
        with Image.open(self.image) as image:
            ImageDraw.Draw(image).text((40, 50), 'Heading', font=ImageFont.truetype(metrics.paths[False, False], 24), fill='black')
            image.save(self.image)
        xml, _ = self.convert(response(elements=[
            element('Title', 'Heading'),
            element('Text', 'مملكة إسبانيا العربية Spain', (30, 200, 550, 300)),
        ]))
        self.assertTrue(xml.xpath('//w:bidi', namespaces=NS))
        self.assertTrue(xml.xpath('//w:rtl', namespaces=NS))
        self.assertTrue(xml.xpath('//w:color[@w:val="000000"]', namespaces=NS))

    def test_font_size_accounts_for_letters_without_descenders(self):
        metrics = FontMetrics()
        font = ImageFont.truetype(metrics.paths[False, False], 28)
        for text in ['Normal text here', 'Carrying on at twelve']:
            box = font.getbbox(text)
            image = Image.new('RGB', (box[2] + 4, box[3] - box[1] + 4), 'white')
            ImageDraw.Draw(image).text((2, 2 - box[1]), text, font=font, fill='black')
            size, _ = source_font_size(image, 0.5, metrics, text)
            self.assertAlmostEqual(size, 14, delta=1)

    def test_borderless_table_uses_centered_headers_and_body_gaps(self):
        from word_export.layout import table_geometry
        image = Image.new('RGB', (900, 400), '#cecece')
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype(FontMetrics().paths[False, False], 16)
        columns = [0, 65, 200, 650, 780, 900]
        for index, heading in enumerate(['No', 'Date', 'Description', 'Credit', 'Balance']):
            center = (columns[index] + columns[index + 1]) / 2
            draw.text((center, 8), heading, font=font, fill='black', anchor='mt')
        for row in range(1, 12):
            for index, text in enumerate([str(row), '2026/10/04', 'A longer transaction description', '20.00', '1,000.00']):
                draw.text((columns[index] + 6, 8 + row * 32), text, font=font, fill='black')
        layout = table_geometry(image, 5, 12)
        self.assertIsNotNone(layout.columns)
        widths = [b - a for a, b in zip(layout.columns, layout.columns[1:])]
        self.assertGreater(widths[2], widths[0] * 3)
        self.assertIsNone(table_geometry(image, 7, 12).columns)

    def test_photographed_small_text_and_long_paragraph_measurements(self):
        metrics = FontMetrics()
        image = Image.new('RGB', (400, 30), '#c8c8c8')
        font = ImageFont.truetype(metrics.paths[False, False], 12)
        ImageDraw.Draw(image).text((4, 4), 'Account date 2026/10/04', font=font, fill='#202020')
        size, _ = source_font_size(image, 0.4, metrics, 'Account date 2026/10/04')
        self.assertGreaterEqual(size, 4)
        self.assertLess(size, 6)
        # Skew can join multiple paragraph baselines into one projection band.
        # Do not fit a whole recognized paragraph to the width of one line.
        image = Image.new('RGB', (400, 80), 'white')
        ImageDraw.Draw(image).text((2, 2), 'Large glyphs', font=ImageFont.truetype(metrics.paths[False, False], 35), fill='black')
        size, _ = source_font_size(image, 0.5, metrics, 'Recognized paragraph with many words. ' * 12)
        self.assertGreater(size, 10)

    def test_uncertain_table_preserves_cells_and_warns_about_ragged_rows(self):
        table = '<table><tr><td>Date</td><td>Very long transaction description with details</td></tr><tr><td>Total</td><td>123.45</td><td>0.00</td></tr></table>'
        xml, warnings = self.convert(response(elements=[element('Table', table, (20, 50, 580, 400))]))
        widths = [int(v) for v in xml.xpath('//w:gridCol/@w:w', namespaces=NS)]
        self.assertGreater(widths[1], widths[0] * 2)
        self.assertTrue(any('inconsistent OCR table row widths' in w for w in warnings))
        self.assertEqual(xml.xpath('//w:tr//w:t/text()', namespaces=NS),
                         ['Date', 'Very long transaction description with details', 'Total', '123.45', '0.00'])

    def test_paper_shadows_do_not_become_graphic_fragments(self):
        image = Image.new('RGB', (600, 800), '#dedede')
        draw = ImageDraw.Draw(image)
        draw.rectangle((50, 100, 120, 200), fill='#c6c6c6')
        draw.rectangle((200, 100, 280, 200), fill='blue')
        boxes, _ = missing_graphics(image, [])
        self.assertFalse(any(a < 100 and c > 60 for a, _, c, _ in boxes))
        self.assertTrue(any(a <= 200 and c >= 280 for a, _, c, _ in boxes))

    def test_cli(self):
        results = self.root / "results.json"
        results.write_text(json.dumps([{"index": 0, "response": response(), "error": None}]))
        command = [sys.executable, "-m", "word_export", "--results", str(results),
                   "--output", str(self.output), str(self.image)]
        run = subprocess.run(command, capture_output=True, text=True, check=False)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertIn("Saved", run.stdout)
        original = self.output.read_bytes()
        results.write_text(json.dumps(response(100, 200)))
        run = subprocess.run(command, capture_output=True, text=True, check=False)
        self.assertEqual(run.returncode, 1)
        self.assertIn("dimensions", run.stderr)
        self.assertEqual(self.output.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
